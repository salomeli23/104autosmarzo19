"""Reportes en Excel (módulo Reportes).

Cada reporte es una lista de hojas: {title, subtitle, columns: [(encabezado, clave, tipo, ancho)], rows: [dict], totals: [claves]}.
Tipos de columna: text | money | int | number | date | datetime.
"""
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLOMBIA_TZ = timezone(timedelta(hours=-5))
HEADER_FILL = PatternFill("solid", fgColor="0F2D46")
TITLE_FONT = Font(bold=True, size=14, color="0F2D46")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="E2E8F0")
THIN = Side(style="thin", color="CBD5E1")
FORMATS = {"money": '"$"#,##0', "int": "#,##0", "number": "#,##0.00", "date": "DD/MM/YYYY", "datetime": "DD/MM/YYYY HH:MM"}

STATUS_LABELS = {
    "agendado": "Agendado", "ingresado": "Ingresado", "en_corte": "En corte", "en_revision": "En revisión",
    "asignado": "Asignado", "en_proceso": "En proceso", "en_calidad": "En calidad",
    "servicio_completado": "Servicio completado", "terminado_falta_pago": "Terminado - falta pago",
    "terminado_pagado": "Terminado - pagado", "entregado": "Entregado", "desistimiento": "Desistimiento",
}
CHARGE_LABELS = {"por_cobrar": "Sin cobrar", "pendiente": "Por revisar", "aprobado": "Aprobado (por pagar)",
                 "rechazado": "Rechazado", "pagado": "Pagado"}
PAYMENT_STATUS_LABELS = {"paid": "Pagado", "partial": "Abono parcial", "pending": "Sin pagar"}
METHOD_LABELS = {"efectivo": "Efectivo", "breb": "Bre-B", "nequi": "Nequi", "daviplata": "Daviplata",
                 "bancolombia": "Bancolombia", "tarjeta_debito": "Tarjeta débito", "tarjeta_credito": "Tarjeta crédito",
                 "transferencia": "Transferencia", "tarjeta": "Tarjeta", "otro": "Otro", "sin_pago": "Sin pago"}


def local_dt(iso: Optional[str]) -> Optional[datetime]:
    """ISO (UTC) → fecha/hora de Colombia sin zona, para Excel."""
    if not iso:
        return None
    try:
        dt = datetime.fromisoformat(str(iso).replace("Z", "+00:00"))
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(COLOMBIA_TZ).replace(tzinfo=None)


def range_bounds(date_from: str, date_to: str) -> tuple:
    """Fechas locales (AAAA-MM-DD) → límites ISO en UTC para comparar con las fechas guardadas."""
    start = datetime.strptime(date_from, "%Y-%m-%d").replace(tzinfo=COLOMBIA_TZ)
    end = datetime.strptime(date_to, "%Y-%m-%d").replace(tzinfo=COLOMBIA_TZ) + timedelta(days=1)
    return start.astimezone(timezone.utc).isoformat(), end.astimezone(timezone.utc).isoformat()


def build_workbook(report_title: str, period: str, sheets: List[Dict[str, Any]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    generated = datetime.now(COLOMBIA_TZ).strftime("%d/%m/%Y %H:%M")
    for sheet in sheets:
        ws = wb.create_sheet(sheet["title"][:31])
        cols = sheet["columns"]
        ws.cell(row=1, column=1, value=f"PolarizadosYA! · {report_title}").font = TITLE_FONT
        ws.cell(row=2, column=1, value=f"{sheet.get('subtitle') or sheet['title']} · {period} · generado {generated}").font = Font(italic=True, color="64748B")
        header_row = 4
        for j, (header, _key, _kind, width) in enumerate(cols, 1):
            c = ws.cell(row=header_row, column=j, value=header)
            c.font, c.fill = HEADER_FONT, HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            ws.column_dimensions[get_column_letter(j)].width = width
        rows = sheet["rows"]
        for i, row in enumerate(rows, header_row + 1):
            for j, (_h, key, kind, _w) in enumerate(cols, 1):
                value = row.get(key)
                if kind in ("date", "datetime") and isinstance(value, str):
                    value = local_dt(value) if "T" in value else (datetime.strptime(value, "%Y-%m-%d") if value else None)
                    if kind == "date" and value:
                        value = value.date()
                c = ws.cell(row=i, column=j, value=value)
                if kind in FORMATS:
                    c.number_format = FORMATS[kind]
                if kind == "text":
                    c.alignment = Alignment(vertical="top", wrap_text=bool(value and len(str(value)) > 40))
                c.border = Border(bottom=THIN)
        last = header_row + len(rows)
        if rows:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{last}"
        totals = sheet.get("totals") or []
        if totals and rows:
            t = last + 1
            ws.cell(row=t, column=1, value="TOTAL").font = Font(bold=True)
            for j, (_h, key, kind, _w) in enumerate(cols, 1):
                ws.cell(row=t, column=j).fill = TOTAL_FILL
                if key in totals:
                    letter = get_column_letter(j)
                    c = ws.cell(row=t, column=j, value=f"=SUBTOTAL(9,{letter}{header_row + 1}:{letter}{last})")
                    c.font = Font(bold=True)
                    c.number_format = FORMATS.get(kind, "#,##0")
        if not rows:
            ws.cell(row=header_row + 1, column=1, value="Sin datos en este periodo").font = Font(italic=True, color="64748B")
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def summarize(rows: List[dict], key: str, fields: Dict[str, str]) -> List[dict]:
    """Agrupa filas por `key` y suma (sum:campo) o cuenta (count / count:campo=valor)."""
    out: Dict[str, dict] = {}
    for r in rows:
        k = r.get(key) or "Sin asignar"
        acc = out.setdefault(k, {key: k, **{f: 0 for f in fields}})
        for f, how in fields.items():
            if how == "count":
                acc[f] += 1
            elif how.startswith("sum:"):
                acc[f] += float(r.get(how[4:]) or 0)
            elif how.startswith("count:"):
                field, _, val = how[6:].partition("=")
                acc[f] += 1 if str(r.get(field)) == val else 0
    return sorted(out.values(), key=lambda x: str(x[key]).lower())
